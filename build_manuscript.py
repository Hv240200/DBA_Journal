"""Reproducible DOCX manuscript builder from validated experiment artifacts."""

from pathlib import Path
import json
import re
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).parent
EXPANDED = '--expanded' in sys.argv
data = json.loads((BASE / 'results' / 'results.json').read_text(encoding='utf-8'))
models = data['model_metrics']
selective = data['retrospective_equal_coverage']
delta = data['paired_bootstrap_delta_hybrid_minus_baseline_at_80pct']
op = data['validation_threshold_operating_point']
audit = data['agreement_audit_heldout']
gate = data['validation_evidence_gate']

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.73)
sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.83)
sec.right_margin = Inches(.83)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0,0,0)
normal.paragraph_format.line_spacing = 2.0
normal.paragraph_format.space_after = Pt(4)

for name, size, before, after in [('SectionHeading',10.8,8,3),('Subheading',10.5,5,1)]:
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.keep_with_next = True

caption = doc.styles.add_style('FigureCaption', WD_STYLE_TYPE.PARAGRAPH)
caption.font.name = 'Times New Roman'
caption.font.size = Pt(8.8)
caption.font.color.rgb = RGBColor(0,0,0)
caption.paragraph_format.line_spacing = 1.1
caption.paragraph_format.space_after = Pt(5)
caption.paragraph_format.keep_with_next = False

def text(value, indent=False):
    p = doc.add_paragraph(value)
    if indent:
        p.paragraph_format.first_line_indent = Inches(.22)
    return p

def section(value):
    # Match the numbered section-heading convention in the accepted KBS example.
    value = re.sub(r'^(\d+)\s+', r'\1. ', value)
    return doc.add_paragraph(value, style='SectionHeading')
def sub(value): return doc.add_paragraph(value, style='Subheading')
def equation(value, n):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.5
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(4)
    r=p.add_run(f'{value}    ({n})')
    r.font.name='Cambria Math'
    r.font.size=Pt(10)

def add_table(headers, rows, widths, caption_text):
    p=doc.add_paragraph(caption_text, style='FigureCaption')
    p.paragraph_format.keep_with_next = True
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for row in rows:
        cells=t.add_row().cells
        for j, v in enumerate(row): cells[j].text = str(v)
    header_properties = t.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement('w:tblHeader')
    repeat_header.set(qn('w:val'), 'true')
    header_properties.append(repeat_header)
    for ridx, row in enumerate(t.rows):
        for j, c in enumerate(row.cells):
            c.width = Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=tcPr.first_child_found_in('w:tcBorders')
            if borders is None:
                borders=OxmlElement('w:tcBorders')
                tcPr.append(borders)
            for edge in ['top','left','bottom','right','insideH','insideV']:
                b=OxmlElement('w:'+edge)
                b.set(qn('w:val'),'single')
                b.set(qn('w:sz'),'4')
                b.set(qn('w:color'),'D9D9D9')
                borders.append(b)
            shading=OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E5EAF0' if ridx==0 else ('F5F7FA' if ridx%2==0 else 'FFFFFF'))
            tcPr.append(shading)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.05
                p.paragraph_format.space_after=Pt(1)
                for r in p.runs:
                    r.font.name='Times New Roman'
                    r.font.size=Pt(8.7)
                    r.font.bold = (ridx==0)
    doc.add_paragraph('').paragraph_format.space_after=Pt(0)
    return t

def figure(filename, caption_text, width):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.keep_with_next=True
    picture = p.add_run().add_picture(str(BASE/'results'/filename), width=Inches(width))
    alt_texts = {
        'figure_1_architecture.png': 'Flowchart: steel measurements feed calibrated ML probabilities and train-derived decision rules. Candidate ranking scores are audited on validation data before automatic labeling or human review.',
        'figure_2_risk_coverage.png': 'Line plot of selective classification error against automatic coverage. At 80 percent coverage, confidence-only error is 14.79 percent and knowledge-guided error is 16.72 percent.'
    }
    picture._inline.docPr.set('descr', alt_texts[filename])
    doc.add_paragraph(caption_text,style='FigureCaption')

title=doc.add_paragraph(style='Title')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.line_spacing=1.1
title.paragraph_format.space_after=Pt(6)
run=title.add_run('Knowledge-Based AI and Data-Driven Decision Support for Digital Transformation, Innovation, and Entrepreneurship in Manufacturing SMEs')
run.font.name='Times New Roman'
run.font.bold=True
run.font.size=Pt(12.5)
run.font.color.rgb=RGBColor(0,0,0)
title.style.font.color.rgb=RGBColor(0,0,0)
style_ppr=title.style._element.get_or_add_pPr()
for border in style_ppr.findall(qn('w:pBdr')):
    style_ppr.remove(border)
direct_ppr=title._p.get_or_add_pPr()
for border in direct_ppr.findall(qn('w:pBdr')):
    direct_ppr.remove(border)

article_type=doc.add_paragraph('Expanded research manuscript — not the Short Communication submission version' if EXPANDED else 'Short Communication')
article_type.alignment=WD_ALIGN_PARAGRAPH.CENTER
article_type.paragraph_format.line_spacing=1.0
article_type.paragraph_format.space_after=Pt(10)
for r in article_type.runs:
    r.font.size=Pt(10)
    r.font.italic=True

section('Abstract')
text('Manufacturing decision support increasingly combines machine-learning predictions with readable knowledge rules, but a plausible rule is not necessarily useful for controlling automation risk. We present a compact, evidence-gated hybrid architecture and investigate whether training-derived decision rules add reliable information beyond calibrated confidence for steel-defect quality triage. Using 1,941 UCI Steel Plates Faults observations, an internally calibrated gradient-boosting model predicts seven classes while a shallow decision tree provides a transparent agreement signal. A stratified 60/20/20 train/validation/test protocol compares confidence-only and knowledge-guided review policies at identical automation budgets. The classifier obtained 77.89% held-out accuracy and 77.76% macro-F1. Although ML–rule disagreement coincided with higher error (32.12% versus 16.67% on agreement cases), the naive hybrid increased observed selective error at 80% coverage from 14.79% to 16.72%; the paired-bootstrap interval for its 1.93-point difference included zero. A validation-only evidence gate rejected the lower-performing candidate on validation. This reproducible negative finding clarifies when symbolic guidance should be audited before deployment; manufacturing-SME transformation and entrepreneurial value are potential applications, not measured effects.')
p=doc.add_paragraph()
p.paragraph_format.line_spacing=1.2
p.add_run('Keywords: ').bold=True
p.add_run('knowledge-based AI; data-driven decision support; selective classification; manufacturing quality; digital transformation; innovation and entrepreneurship')

section('1 Introduction')
text('Data-intensive manufacturing promises better inspection decisions, yet high-quality process data and organizational readiness are persistent barriers, particularly in SMEs [6–9]. Machine learning can classify observed faults, but plant personnel also need a defensible policy for deciding which predictions can be accepted automatically and which require manual review. Selective prediction formalizes this trade-off as the error among accepted cases at a specified coverage [2].')
text('Knowledge-based AI promises interpretability through explicit decision rules. However, an intelligible rule and a reliable rule are different things: a rule extracted from the same limited data as the predictor can be wrong or poorly aligned with prediction confidence. We therefore ask whether a shallow induced knowledge model contributes information useful for selecting safe automatic quality classifications beyond calibrated machine-learning confidence. The narrow experiment focuses on quality triage, not physical corrective actions, production optimization, or a demonstrated business transformation.')
text('This distinction is consequential for manufacturing SMEs: digital-transformation studies identify technological, managerial, implementation, and financing constraints [23–27]. An inspectable rule layer may be attractive to quality engineers; nevertheless, a conservative-looking rule can move correct high-confidence cases into review while allowing less reliable cases through. We test whether including such rules improves the selective-risk/coverage frontier under a fixed human-review budget.')

section('2 Related Work and Contribution')
text('Probability calibration [3] and selective classification [2] motivate confidence-sensitive review. Industrial studies have examined uncertainty-aware evidential reasoning [5], reliable learning with sparse data [13], data–knowledge synergy [4], and knowledge-graph fault diagnosis [12,14,15,18]. Other machine-fault studies address noisy labels, limited labeled examples, and data-driven diagnostic methods [16,17,19,20]. Hybrid reasoning is therefore established; combining a tree with a classifier is not itself novel.')
text('Related steel-surface studies emphasize visual defect recognition and neural feature extraction [21,22], while human-in-the-loop industrial inspection supports explainable operator interaction [10]. The specific gap examined here is methodological: rule agreement can correlate with prediction failure yet contain no useful information beyond confidence. Aggregate classification accuracy cannot establish a rule layer\'s value as a review-allocation policy.')
text('We evaluate an evidence gate: (i) construct an agreement-based selection candidate and (ii) activate it only if a separate validation sample shows lower selective error than confidence-only selection at a fixed budget. We audit both policies on an untouched test set with paired uncertainty. This is a reproducible, narrowly scoped decision-policy test, not a claimed invention of a new hybrid-learning algorithm.')
if EXPANDED:
    sub('2.1 Manufacturing data as decision evidence')
    text('The industrial context matters because the same numerical model can support substantially different decisions. A quality engineer may require a ranked list for secondary inspection, whereas a production supervisor may need a release recommendation, and a process engineer may investigate whether recurring measurements indicate equipment drift. Accuracy summarizes predictive classification on a labeled population but does not say whether uncertain instances reach the personnel equipped to resolve them. The industrial-data literature emphasizes that a useful analytics project requires an adequate representation of processes, reliable data acquisition and agreement between domain experts and data scientists about what each variable describes [7]. Consequently, a classification benchmark supports only the narrow data-to-label portion of decision intelligence. Evidence about disposition actions, operator capacities and economic costs is necessary before the broader factory-level decision can be evaluated.')
    text('SME research provides a second reason to distinguish model performance from an operational decision. Adoption constraints are not only technical. Resource allocation, organizational change, workforce capabilities and business-model alignment influence whether a promising AI pilot becomes a sustainable activity [6,8,23]. Research on manufacturing firms also links effective data-driven decision-making to data integration and digital technologies [9]. These findings motivate auditable, low-friction deployments; they do not establish that any particular defect classifier will reduce SME costs. The proposed protocol therefore uses manufacturing quality triage as an experimentally tractable component and treats digital transformation and entrepreneurship as a hypothesis about future adoption, not an inferred result from the UCI records.')
    sub('2.2 Distinguishing symbolic knowledge from agreement heuristics')
    text('Industrial knowledge can enter an AI system through typed entities and relationships, operating constraints, cause–effect mechanisms, authored procedures, exception lists, and expert judgments about uncertainty. Graph-based fault diagnosis illustrates a richer family of knowledge representation, in which equipment relationships or multimodal evidence mediate diagnostic inference [12,14,15,18]. Evidential methods can explicitly represent incomplete or unfamiliar fault evidence [5], while integrated data–knowledge systems address few-shot adaptation [4]. Their objectives and evaluation settings differ from a depth-four tree trained only to reproduce existing labels. For clarity, we refer to the latter as a knowledge-style symbolic proxy; its branch conditions are interpretable patterns learned from data, not independently validated industrial facts.')
    text('This distinction identifies a potentially important failure mode. If two predictors are trained on essentially the same examples and respond to similar feature variation, their agreement may only restate whether the underlying case is easy. Even where disagreement is associated with a larger raw error rate, a confidence score may already encode much of that difficulty. The marginal value of the symbolic signal should therefore be judged conditional on the score used by the real decision policy. Merely displaying a consistent second label or a plausible path explanation does not prove that a different set of cases should be released without inspection. This is particularly important where an interpretable rule tree is less accurate than the primary predictor.')
    sub('2.3 Position against existing quality-inspection systems')
    text('Steel-defect analysis spans sensor-derived numerical measurements and direct imaging. Visual architectures such as the ST-YOLO and DSTEELNet models are designed to detect or classify surface patterns in images [21,22]. They help establish the industrial relevance of steel-surface defect recognition but are not like-for-like predictive baselines for the tabular features available in the present benchmark. The related question of operator involvement has been explored in explainable, human-in-the-loop hard-metal inspection systems [10]. Our analysis isolates one upstream policy choice: whether an independently readable rule recommendation adds information to the ranking of test instances proposed for automatic classification. It neither evaluates image localization nor reproduces an interactive industrial platform.')

section('3 Data and Method')
sub('3.1 Dataset and partition')
text('The Steel Plates Faults dataset [1] contains 1,941 observations with 27 numerical or binary inputs and one of seven fault labels. It contains no missing values according to its repository metadata. The archived matrix was checked for its expected 34 columns and one-hot target integrity. A fixed stratified split yielded 1,164 training, 388 validation, and 389 held-out test examples (seed 42). No test labels were used to fit models, induce rules, calibrate probabilities, or select operational cutoffs.')
text('These are industrial fault records, not SME adoption observations; the source manufacturer is not assumed to be an SME. The unmodified upstream archive is identified by a SHA-256 digest. Manufacturing data-foundation methodologies [7] support treating provenance as a prerequisite to deployment rather than a peripheral documentation concern.')
sub('3.2 Probabilistic and symbolic models')
text('A histogram-based gradient-boosted tree ensemble (180 boosting iterations, 15 maximum leaves, learning rate 0.07) produces multiclass probabilities. Sigmoid calibration uses three-fold cross-validation inside the training partition. Logistic regression with standardized inputs and a standalone shallow decision tree serve as predictive baselines. The symbolic component is a separate depth-four decision tree with a minimum of 30 samples per leaf. Its decision paths are exported as legible IF-THEN rules; these are induced from training data, not elicited from manufacturing experts or validated SOPs.')
text('Let f(x) be the calibrated class-probability vector and g(x) the label assigned by the independent rule tree. The selected label is the maximum-probability class and c(x) is its probability. The agreement indicator A(x) equals one when the two predicted labels match. The fixed candidate knowledge-guided score applies a 35% confidence penalty to disagreement cases; this parameter is an illustrative predeclared intervention, not an optimized hyperparameter.')
equation('ŷ(x) = arg maxₖ fₖ(x);  c(x) = maxₖ fₖ(x);  A(x) = 1[g(x) = ŷ(x)]',1)
equation('s₀(x) = c(x);    sK(x) = c(x)[0.65 + 0.35A(x)]',2)
sub('3.3 Validation-evidence gate')
text('The proposed reliability gate tests whether the knowledge-guided ranking deserves to control decisions. At 80% target automatic coverage, validation records are ranked once by s₀ and once by sK; the lower observed selective-error policy is selected, with confidence-only winning a tie. The selected policy is then frozen before test evaluation. This safeguard does not guarantee improvement in new factories, but it prevents the symbolic component from being accepted merely because its logic appears plausible.')
equation('p* = 0 if Rvalid(s₀, 0.80) ≤ Rvalid(sK, 0.80); K otherwise',3)
text('A deployment would also require validation-derived cutoffs, process-owner approval of rule semantics, monitoring, and adjudication of rejected cases. Figure 1 distinguishes fitting a symbolic proxy from authorizing its use; rejection of the proposed hybrid is a valid gate outcome.')
figure('figure_1_architecture.png','Figure 1. Evidence-gated knowledge-based decision-support architecture. Induced decision paths are inspectable but not expert-approved SOPs. The validation comparison controls whether the rule-assisted candidate is used.',6.25)
if EXPANDED:
    sub('3.4 Provenance and interpretation of the rule layer')
    text('An explicit separation is maintained among three distinct objects: the calibrated predictive distribution, the induced decision paths, and the operational acceptance policy. The first estimates class probabilities given measured inputs; the second offers concise statements about training-set partitions; the third assigns scarce review attention. Although the decision tree can be represented by IF–THEN branches, that textual form alone does not justify calling its statements shop-floor knowledge. An expert-approved knowledge base would require authorship, provenance, rule scope, versioned source documents, an exception mechanism, and verification against plant safety and quality obligations. None of those external endorsements is present in the UCI data, and none is inferred by the algorithm.')
    text('The hybrid score deliberately leaves the ensemble class label unchanged. Under agreement, a case retains confidence c(x); under disagreement, the score is multiplied by 0.65. The multiplier affects ranking and the resulting selection decision but does not recalibrate the underlying class probabilities. Treating it as a probability would be incorrect because it is a policy score rather than a calibrated probability estimate. At equal coverage, the evaluation asks whether the score improves the ordering of accepted versus rejected cases. The design therefore avoids conflating a modest change in review assignments with a classification accuracy gain. A model that changes only its review ranking must be assessed with selective risk and coverage, not by claiming that its base classifier learned a better decision boundary.')
    sub('3.5 Dataset provenance, isolation and validation responsibilities')
    text('The training/validation/test split is stratified so that every subset retains representation from each defect class. Training data serve model fitting, tree induction and the internal three-fold calibration folds. The distinct validation subset is used for the policy comparison at the predeclared 80% automatic-coverage budget and for score-cutoff estimation. The held-out test subset is reserved for reporting predictive baselines, matched-coverage comparisons, agreement diagnostics and paired uncertainty. The analysis does not use held-out labels to change the score penalty, the model architecture or the selected policy. This is the essential division of labor in an experiment where a second signal can influence automation decisions without ever changing predicted classes.')
    text('The UCI source ZIP and its recorded digest establish which historical observations underlie the calculations. The implementation also checks the expected numeric data matrix and the integrity of the one-hot encoded class labels. No bespoke cleaning pipeline, missing-value imputation or factory-specific preprocessing is claimed. A deployment on contemporary line data would have to document sensor calibration, unit conventions, defect label ontology, sampling cadence and missingness patterns before using the same code. This distinction is necessary because exact computational reproducibility on a public benchmark is a weaker claim than measurement validity in a production environment.')
    sub('3.6 From classification error to action-dependent loss')
    text('Selective risk treats every wrong accepted label equally. In a real manufacturing setting, errors may carry different consequences: a misclassified minor cosmetic anomaly and a missed safety-critical defect need not receive the same penalty. The downstream objective would therefore incorporate a cost matrix indexed by true defect, released decision and review action, plus the time or capacity required for inspection. One can then compare expected decision loss between policies subject to a maximum inspection workload. These quantities are not contained in the benchmark, so such a cost-sensitive objective cannot be numerically evaluated here. It is specified as a requirement for a future operational study rather than silently assumed by the accuracy and macro-F1 measures.')

section('4 Experimental Protocol')
text('Predictive baselines are evaluated using held-out accuracy, macro-F1, multiclass Brier score, and log loss. Multiclass Brier score averages the per-case sum of squared differences between predicted probabilities and one-hot target indicators. This proper probabilistic metric complements class-label accuracy because the review policy depends on probabilities rather than just categorical outputs.')
text('For ranking score s, let S(s,q) contain the top q-fraction of scores. Selective risk R(s,q) is the error proportion in this subset. Matched subsets at q = 0.60, 0.80, and 0.90 isolate ranking quality because the policies share identical base predictions. These test-set rankings are retrospective diagnostics, never live threshold-selection rules.')
equation('R(s,q) = [ Σᵢ∈S(s,q) 1(ŷᵢ ≠ yᵢ) ] / |S(s,q)|',4)
text('Separately, the top-80%-validation score cutoff for each policy is frozen and applied to test examples, so observed operational coverage can drift. A 2,000-resample paired test-set bootstrap yields a percentile 95% interval for R(sK,0.80) − R(s₀,0.80), recomputing matched ranks in each resample. It quantifies conditional sample variation, not model-retraining uncertainty.')
if EXPANDED:
    sub('4.1 Predictive and decision-policy baselines')
    text('The logistic-regression and shallow-tree classifiers test two straightforward predictive alternatives with different inductive biases. The primary classifier is a histogram-based gradient-boosting ensemble with sigmoid calibration, all fitted without validation or test labels. These baselines provide a narrow sanity check: the study would be difficult to interpret if a weak primary predictor were used to justify claims about a knowledge layer while a simpler model performed substantially better. However, only three predictive approaches are included; the experiment does not exhaust neural networks, anomaly detectors, class-imbalance methods, probabilistic reasoning models or industrial-domain transfer learning. Accordingly, the result is framed as a policy-selection comparison around one fixed primary classifier rather than as a general model-leaderboard result.')
    text('The two principal policies share precisely the same multiclass predictions. A confidence-only selector ranks records using maximum calibrated probability. A knowledge-guided selector reduces the rank of predictions that disagree with the shallow rule tree by an explicitly fixed factor. This controls a key confounder: differences in the accepted test-set errors can be attributed to which records are selected, rather than to altered defect labels. Equal-coverage comparisons are especially important because any policy can appear safer by sending more examples to human review. Evaluating the same number of accepted instances at each budget makes this trade-off directly observable.')
    sub('4.2 Fixed operating thresholds versus retrospective ranking curves')
    text('A retrospective curve is informative for diagnosis because a researcher can sort an already observed labeled evaluation sample and ask how many errors would remain among its highest-scoring records. It does not describe a deployable process: a production stream cannot assume that a known fraction of future cases will have high enough scores, nor can it use their unavailable ground-truth labels. The operational counterpart requires an acceptance threshold estimated from data set aside before test evaluation. That is why the study separately stores the validation-score cutoff for the nominal 80% target and applies it unchanged to test scores. Deviations in realized test coverage are expected when score distributions vary and are not recalibrated away using held-out information.')
    text('The validation-evidence gate and the final policy are also distinct from the retrospective plotted comparison. The gate uses validation labels to decide which candidate ranking policy is preferred at the declared budget, with baseline retention on a tie. This rule is a simple form of evidence-based model selection, not a statistical guarantee of improvement. If many penalties, thresholds, rule structures or coverage targets were explored and only the best validation outcome reported, selection bias could arise. The current candidate set is intentionally restricted. A larger experiment should prespecify this search space, retain a final untouched external test source, and investigate whether the gate remains reliable across multiple factories and production periods.')
    sub('4.3 Uncertainty, inferential scope and reproducibility')
    text('The paired percentile bootstrap resamples whole held-out observations and recomputes the accepted subsets for each ranking policy at the same coverage. Pairing is appropriate because the competing decisions concern the same records and the same ground-truth labels. The resulting interval estimates variation associated with that finite evaluation sample conditional on its fitted models. It does not account for different random initializations, a new train–test split, changing class frequencies, newly acquired expert rules or concept drift. Therefore, an interval spanning zero rules out a strong claim that the observed positive risk difference demonstrates actual degradation in every deployment setting; the correct conclusion is that the candidate failed to show an improvement here.')
    text('Reproduction uses the archived UCI input, explicit random seed, fixed model parameters, versioned Python dependencies and a program that writes both visualizations, the exported symbolic rules and a structured results file. A separate rerun from the unchanged source ZIP produced a byte-identical numerical JSON artifact and identical figure checksums in the editorial audit. Such exactness helps reviewers locate coding errors or verify table transcription. It does not establish that the proposed knowledge layer would be trustworthy under changing operating conditions, where monitoring and periodic domain-expert review remain mandatory.')

section('5 Results')
table1=[]
for label in ('Logistic regression','Induced rules (tree)','Calibrated gradient boosting'):
    m=models[label]
    table1.append((label,f'{100*m["accuracy"]:.2f}',f'{100*m["macro_f1"]:.2f}',f'{m["brier"]:.4f}'))
add_table(['Model','Accuracy (%)','Macro-F1 (%)','Brier score'],table1,[3.08,1.12,1.24,1.1],
          'Table 1. Held-out defect classification (n = 389). Lower Brier scores indicate better probabilistic accuracy.')
text(f'The calibrated gradient-boosted classifier was the strongest tested predictive model, attaining {100*models["Calibrated gradient boosting"]["accuracy"]:.2f}% accuracy and {100*models["Calibrated gradient boosting"]["macro_f1"]:.2f}% macro-F1 (Table 1). The rule tree agreed with its predicted class on {100*data["rule_agreement_heldout"]:.2f}% of held-out examples. Agreement alone did not provide a more reliable acceptance ordering.')
text(f'The agreement audit exposes an instructive contrast. Among {audit["agree"]["n"]} agreement cases, the ensemble error rate was {100*audit["agree"]["error_rate"]:.2f}% and mean ensemble confidence was {100*audit["agree"]["mean_ml_confidence"]:.2f}%. Among {audit["disagree"]["n"]} disagreement cases, error was markedly higher at {100*audit["disagree"]["error_rate"]:.2f}%, but mean confidence was already lower at {100*audit["disagree"]["mean_ml_confidence"]:.2f}%. Thus, disagreement correlates with difficulty, yet this observation does not show incremental decision value after confidence is considered. A hybrid system must establish the latter empirically.')
table2=[]
for k in ('0.6','0.8','0.9'):
    v=selective[k]
    b=100*v['Confidence only']['error']
    h=100*v['Knowledge guided']['error']
    table2.append((f'{100*float(k):.0f}% ({v["Confidence only"]["accepted"]})',f'{b:.2f}',f'{h:.2f}',f'{h-b:+.2f}'))
add_table(['Coverage (accepted n)','Confidence only (%)','Knowledge guided (%)','Difference (pp)'],table2,[1.66,1.67,1.67,1.54],
          'Table 2. Retrospective selective error at matched coverage. Positive differences favor confidence-only selection.')
text(f'At 80% automated coverage, confidence-only selection had a selective error of {100*selective["0.8"]["Confidence only"]["error"]:.2f}%, compared with {100*selective["0.8"]["Knowledge guided"]["error"]:.2f}% for the hybrid (Table 2). The difference was {100*delta["point"]:+.2f} percentage points, with a paired-bootstrap 95% interval from {100*delta["ci95_lower"]:+.2f} to {100*delta["ci95_upper"]:+.2f} points. The interval spans zero, so the point estimate should not be described as conclusive evidence of degradation. At 60% coverage, the hybrid ranked 12.02% incorrect cases among accepted records versus 7.73% for confidence alone; the difference narrowed by 90% coverage. Figure 2 displays diagnostic risk–coverage curves.')
figure('figure_2_risk_coverage.png','Figure 2. Held-out risk-coverage curves. The plotted fractions are retrospective equal-coverage rankings, not pre-deployment thresholds.',5.45)
text(f'Using cutoffs fixed on validation data for a nominal 80% automated-coverage target, confidence-only selection achieved {100*op["Confidence only"]["heldout_coverage"]:.2f}% test coverage with {100*op["Confidence only"]["heldout_selective_error"]:.2f}% selective error. The hybrid achieved {100*op["Knowledge guided"]["heldout_coverage"]:.2f}% coverage with {100*op["Knowledge guided"]["heldout_selective_error"]:.2f}% error. These operational points are not coverage-matched and should not replace the controlled comparison in Table 2.')
text(f'The evidence gate makes a separate prospective decision using only validation labels. Validation selective risk at 80% was {100*gate["validation_selective_error_at_80pct"]["Confidence only"]:.2f}% for confidence-only and {100*gate["validation_selective_error_at_80pct"]["Knowledge guided"]:.2f}% for knowledge-guided selection. Accordingly, the gate chose the {gate["selected_policy"].lower()} policy. Its held-out selective error at equal 80% coverage was {100*gate["heldout_selective_error_at_80pct"]:.2f}%. The evidence gate contributes a conservative decision protocol, not a new predictive accuracy gain; it correctly refrained from activating the candidate rule penalty in this experiment.')
if EXPANDED:
    sub('5.1 What the agreement audit does and does not establish')
    text('The apparent discrepancy between the raw agreement audit and the policy comparison is a central empirical lesson. Agreement partitions the test set into groups of different observed difficulty, but a useful review policy needs to order cases within and across those groups better than the existing confidence score. The audit shows that disagreement cases have both a higher error fraction and lower ensemble confidence. Thus, there is no paradox in observing that the agreement label is associated with error while a ranking penalty based on that label produces worse selective risk. The original predictor already assigns many disagreement cases relatively low scores, and a uniform penalty can displace other cases that were actually handled correctly. Group membership is not proof of incremental predictive information.')
    text('At 60% coverage, the candidate suffers a larger selective-risk deficit than at the higher two plotted budgets. This observation is compatible with the strongest-confidence portion of the queue being particularly sensitive to inappropriate rank reversals. It should not be overinterpreted as a universal trend, because only one test sample and one fixed penalty are available. Equally, the 80% bootstrap interval includes zero, so the paper does not claim statistically established inferiority. The recommendation to retain confidence-only selection is based on the observed validation policy choice, with the independent test sample functioning as an audit rather than a parameter-fitting opportunity.')
    sub('5.2 Why accuracy alone would conceal the selection failure')
    text('All selected cases receive the same ensemble class labels regardless of policy. A conventional model-accuracy summary therefore cannot distinguish the two selection rules. Both are coupled to the same base classifier, yet their accepted subsets differ. Table 1 answers which predictive baselines performed better as classifiers; Table 2 asks which ranking policy leaves fewer errors among cases handled automatically at the same nominal budget. Figure 2 extends that question over a range of retrospective coverage fractions. These outputs are complementary, not competing, because none of them directly measures human inspection costs, defect escape severity or operational benefit to an SME. This separation prevents a common interpretive error in applied AI research: treating gains in an auxiliary technical metric as evidence of improved decisions without specifying how the metric changes the workflow.')

section('6 Discussion and Limitations')
text('The negative result is informative because the additional symbolic model is intuitive but substantially less accurate than the calibrated ensemble. Its agreement signal can suppress high-confidence correct predictions and does not necessarily isolate the ensemble\'s mistakes. The findings support directly testing a rule layer rather than assuming that knowledge-style structure improves decisions. This experiment neither disproves knowledge-based reasoning in general nor evaluates expert-authored constraints, reasoning graphs, or neuro-symbolic models developed for other settings [4,5].')
text('The apparently contradictory results—higher raw error under disagreement, but inferior hybrid selection—have a coherent explanation. The classifier is less confident on disagreement cases even without rule information. Applying the same large multiplier to every disagreement ignores heterogeneity: some correct and high-confidence predictions are moved behind lower-quality agreement cases. The relevant question for symbolic reasoning is therefore conditional incremental information, not raw correlation between agreement and failure. A validation benchmark is an inexpensive first filter for such claims, although its result remains dependent on sample composition and the chosen review budget.' if EXPANDED else 'Disagreement cases already received lower ML confidence, so a uniform hybrid penalty could demote correct cases behind lower-quality agreement cases. The relevant question is whether a symbolic signal improves review ranking beyond confidence, not merely whether disagreement correlates with error. The validation comparison is only a preliminary filter and does not guarantee reliability under other samples or review budgets.')
sub('6.1 Interpretation for digital transformation and innovation')
text('In a hypothetical SME workflow, measurements feed calibrated predictions and rule-path explanations; a validated selector sends uncertain cases to an inspector. Related work has studied manufacturing decision-support allocation [11], human-in-the-loop metal inspection [10], and SME AI resource orchestration [8]. A low-cost audit service could be an entrepreneurial opportunity, but adoption, commercialization, and economic benefit were not measured here.')
text('A field study should collect SOP provenance, labor time, misclassification severity, false-release events, drift, override reasons, and throughput; management studies identify data integration as consequential for effective firm-level decision-making [9]. Its target should be expected decision loss under a preapproved policy, not accuracy alone.' if EXPANDED else 'A future SME study should measure SOP provenance, inspection costs, defect severity, throughput, and drift. Firm-level research links data integration to manufacturing decision-making [9], but the relevant operational target is action cost rather than accuracy alone.')
sub('6.2 Threats to validity')
text('The benchmark is small and historical, and the experiment uses a single stratified partition. Rules are label-induced proxies, not knowledge-graph assertions [12,14,15,18] or approved SOPs. No action costs, temporal drift, external factories, SME status, or financial outcomes are observed. Bootstrap uncertainty conditions on fitted models; the fixed 0.65 penalty is illustrative, not an optimized policy.')
text('The gate can overfit if many policies are tried on one validation set. Future studies need repeated and out-of-domain splits, expert-authored constraints, prospective review-capacity validation, and real process outcomes. The present negative baseline demonstrates neither superiority to neuro-symbolic alternatives nor digital-entrepreneurship gains.')
if EXPANDED:
    sub('6.3 A governed manufacturing deployment pathway')
    text('A responsible implementation would begin with an inventory of existing inspection tasks, release criteria and documented standard operating procedures. The process engineer would identify which measured features correspond to actual equipment or product characteristics and whether the data can be linked to repeatable production contexts. Quality personnel would define which fault dispositions are advisory, which require mandatory manual inspection and which cannot legally or operationally be automated. Only then should AI predictions be routed through a policy that can be audited against these requirements. A confidence or rule score cannot override a hard safety constraint merely because aggregate statistical performance appears favorable.')
    text('The proposed architecture can be integrated into a digital quality workflow without implying that the benchmark evaluates such integration. After acquiring measurements, a model service returns the predicted defect class and a confidence score. A governed rule service attaches the relevant versioned human-authored constraints, provenance and any known exceptions. An evidence-tested selector either forwards an advisory disposition to an existing quality record or requests human adjudication. The final decision, override reason, elapsed inspection time, author and source-system version should be recorded so that monitoring can distinguish prediction errors from procedural changes. Manufacturing execution systems and KPI programs provide relevant organizational context for that integration [24]; adoption must nevertheless be evaluated in a live environment.')
    text('Monitoring should operate on several levels. Predictive drift concerns changes in the relationship between measured features and fault labels. Data drift concerns upstream sensor signals, missingness, measurement scales or production mix. Knowledge drift occurs when a procedure changes, an engineering deviation is approved or an experienced quality inspector discovers an unmodeled exception. Policy drift arises when incoming score distributions change the fraction of cases accepted under a fixed numerical threshold. Each can degrade performance through a different pathway, so aggregate model accuracy is inadequate as the sole monitoring indicator. A versioned audit trail can support disciplined rollback, scheduled recalibration and an explicit human escalation path.')
    sub('6.4 SME organizational value and entrepreneurial hypotheses')
    text('Resource-limited manufacturers may prefer incremental decision-support projects that can be attached to existing quality processes rather than an immediate enterprise-wide platform replacement. A small, inspectable service has potential advantages in deployment effort, but the evidence needed to evaluate these advantages is organizational and economic. Costs include sensor integration, software maintenance, operator retraining, validation labor, audit preparation, cybersecurity controls and the consequences of mistaken releases. Benefits, if any, would arise from reduced repetitive inspection, faster identification of recurring defects, more consistent triage and improved documentation. The benchmark contains none of these cost or benefit measurements, so a numerical return on investment cannot be inferred.')
    text('The innovation and entrepreneurship components of the title consequently remain a research agenda. An internal innovation pathway could investigate whether the service improves the speed and quality of cross-functional action among quality, production, design, inventory and operations personnel. An external entrepreneurial pathway could investigate whether a configurable, auditable inspection-support product is valued by SMEs that cannot sustain a large data-science function. Published SME digital-transformation work discusses implementation strategies, managerial readiness and financial access [23,25–27], while field-based AI adoption work emphasizes resources and capabilities [8]. Those studies help formulate organizational hypotheses but do not validate a revenue model for the present system.')
    sub('6.5 Research design for a confirmatory multi-site evaluation')
    text('A more discriminating follow-up would compare at least three knowledge conditions: confidence-only selection; the present data-induced tree proxy; and a genuinely independent, expert-authored rule base. A fourth condition could test graph-based constraints or calibrated evidential reasoning if its representation and computational costs are specified in advance. Rules should be recorded with subject-matter authors, dates, relevant process conditions and explicit exceptions. For every comparison, the primary metric would be predeclared expected quality decision loss under a fixed inspection-capacity constraint, followed by selective risk, realized review load and human override frequency. Independent factory lines and later production periods should provide external tests, with uncertainty intervals that incorporate resampling or refitting of the learned models.')
    text('A confirmatory operational pilot would also measure cycle time, defect escape rate, false rejection cost, operator trust, usability and audit effort. Ethical and operational safeguards would include manual review of critical dispositions, stop conditions for unacceptable drift, and documentation of interventions made by people rather than inferred from the algorithm. A design in which neither personnel nor production conditions are observed cannot directly test these consequences. By stating these requirements explicitly, the present research avoids elevating a benchmark classification task into a claim about manufacturing transformation while providing a practical blueprint for assessing whether knowledge-based decision support earns its organizational role.')

section('7 Conclusion')
text(('This expanded study examines an evidence-gated decision protocol for assessing whether a readable symbolic rule layer adds information beyond calibrated machine-learning confidence for manufacturing quality triage. ' if EXPANDED else 'This short communication introduces an evidence-gated protocol for evaluating whether a transparent knowledge layer adds decision value beyond calibrated machine-learning confidence in manufacturing quality triage. ') + 'On the reported steel-fault benchmark, disagreement identified a higher-risk subgroup, but the naive agreement penalty failed to improve risk at matched review budgets. The validation gate rejected that penalty and retained the safer observed baseline. The central lesson is methodological: knowledge-guided decision support should be deployed on demonstrated incremental value rather than on interpretability alone. Transfer to manufacturing SMEs, cost reduction, and new digital services require expert-authored rules and independent operational trials.')

section('Data and Code Availability')
text(f'The Steel Plates Faults dataset is publicly available under CC BY 4.0 [1]. Source archive SHA-256: {data["dataset_sha256"]}. The minimal Python source and reproduction instructions are at https://github.com/Hv240200/DBA_Journal (verified revision: 89bf6b99097017a1cefd5d5d4d3e988a1c1b8c56); generated numerical outputs and both figures accompany this manuscript and are reproducible by running the code. No private factory records were used.')

section('Declaration of Generative AI Assistance')
text('OpenAI ChatGPT assisted with manuscript drafting, research-code development, and figure layout. The human authors must verify the analysis, interpretation, citations, and final wording and remain fully responsible for the submitted work. This draft is not approved for submission without that review.')

sub('Declaration of Competing Interest')
text('Author confirmation required: the authors must supply and approve their actual competing-interest disclosure. No declaration has been presumed.')
sub('Funding')
text('Author confirmation required: list every relevant funding source and grant identifier, or confirm that no external funding supported the work.')
sub('CRediT Authorship Contribution Statement')
text('Author confirmation required: identify all authors and verify each person\'s actual CRediT roles before submission.')

section('References')
references = [
    '[1] M. Buscema, S. Terzi, W. Tastle, Steel Plates Faults, UCI Machine Learning Repository (2010). https://doi.org/10.24432/C5J88N.',
    '[2] Y. Geifman, R. El-Yaniv, Selective classification for deep neural networks, Adv. Neural Inf. Process. Syst. 30 (2017). https://papers.neurips.cc/paper/2017/hash/4a8423d5e91fda00bb7e46540e2b0cf1-Abstract.html.',
    '[3] C. Guo, G. Pleiss, Y. Sun, K.Q. Weinberger, On calibration of modern neural networks, Proc. Int. Conf. Mach. Learn. 70 (2017) 1321–1330. https://proceedings.mlr.press/v70/guo17a.html.',
    '[4] C. Chen, X. Si, Z. Wang, Z. Zhang, C. Hu, DEFINED: A data-knowledge synergistic framework for few-shot class-incremental fault diagnosis, Knowledge-Based Systems 347 (2026) 116153. https://doi.org/10.1016/j.knosys.2026.116153.',
    '[5] H. Fu, Q. Wang, C. Gong, Z.-G. Su, Incremental open-set evidential reasoning for uncertainty-aware fault diagnosis in rotating machinery, Knowledge-Based Systems 344 (2026) 116090. https://doi.org/10.1016/j.knosys.2026.116090.',
    '[6] T. Masood, P. Sonntag, Industry 4.0: Adoption challenges and benefits for SMEs, Computers in Industry 121 (2020) 103261. https://doi.org/10.1016/j.compind.2020.103261.',
    '[7] S. Meier, S. Klarmann, N. Thielen, et al., A process model for systematically setting up the data basis for data-driven projects in manufacturing, Journal of Manufacturing Systems 71 (2023) 1–19. https://doi.org/10.1016/j.jmsy.2023.08.024.',
    '[8] E. Peretz-Andersson, S. Tabares, P. Mikalef, V. Parida, Artificial intelligence implementation in manufacturing SMEs: A resource orchestration approach, International Journal of Information Management 77 (2024) 102781. https://doi.org/10.1016/j.ijinfomgt.2024.102781.',
    '[9] R. Colombari, A. Geuna, S.R. Helper, et al., The interplay between data-driven decision-making and digitalization: A firm-level survey of the Italian and U.S. automotive industries, International Journal of Production Economics 255 (2023) 108718. https://doi.org/10.1016/j.ijpe.2022.108718.',
    '[10] T. Kotsiopoulos, G. Papakostas, T. Vafeiadis, et al., Revolutionizing defect recognition in hard metal industry through AI explainability, human-in-the-loop approaches and cognitive mechanisms, Expert Systems with Applications 255 (2024) 124839. https://doi.org/10.1016/j.eswa.2024.124839.',
    '[11] E. Mahmoodi, M. Fathi, M. Tavana, et al., Data-driven simulation-based decision support system for resource allocation in industry 4.0 and smart manufacturing, Journal of Manufacturing Systems 72 (2024) 287–307. https://doi.org/10.1016/j.jmsy.2023.11.019.',
    '[12] C. Peng, Y. Sheng, W. Gui, et al., Bearing fault diagnosis based on multimodal knowledge graphs under few-shot samples, Knowledge-Based Systems 330 (2025) 114604. https://doi.org/10.1016/j.knosys.2025.114604.',
    '[13] M. Feng, H. Shao, M. Shao, et al., Utilizing Bayesian generalization network for reliable fault diagnosis of machinery with limited data, Knowledge-Based Systems 305 (2024) 112628. https://doi.org/10.1016/j.knosys.2024.112628.',
    '[14] Z. Mao, H.M. Wang, B. Jiang, J. Xu, H. Guo, Graph convolutional neural network for intelligent fault diagnosis of machines via knowledge graph, IEEE Transactions on Industrial Informatics 20 (5) (2024) 7862–7870. https://doi.org/10.1109/TII.2024.3367010.',
    '[15] C. Peng, Y. Sheng, W. Gui, Z. Tang, C. Li, A rolling bearing fault diagnosis method based on multimodal knowledge graph, IEEE Transactions on Industrial Informatics 20 (11) (2024) 13047–13057. https://doi.org/10.1109/TII.2024.3431074.',
    '[16] F. Dunkin, X. Li, Z. Zhang, et al., Certainty from uncertainty: Multigranularity labeling inspired by quantum collapse for learning with noisy labels in fault diagnosis, IEEE Transactions on Industrial Informatics 21 (8) (2025) 6443–6454. https://doi.org/10.1109/TII.2025.3567264.',
    '[17] X. Zhang, Z. Su, X. Hu, Y. Han, S. Wang, Semisupervised momentum prototype network for gearbox fault diagnosis under limited labeled samples, IEEE Transactions on Industrial Informatics 18 (9) (2022) 6203–6213. https://doi.org/10.1109/TII.2022.3154486.',
    '[18] Y. Chen, M.J. Zhou, M. Zhang, M. Zha, Knowledge-graph-driven fault diagnosis methods for intelligent production lines, Sensors 25 (13) (2025) 3912. https://doi.org/10.3390/s25133912.',
    '[19] L. Lei, W. Li, S. Zhang, C. Wu, H. Yu, Research progress on data-driven industrial fault diagnosis methods, Sensors 25 (9) (2025) 2952. https://doi.org/10.3390/s25092952.',
    '[20] S. Qiu, X. Cui, Z. Ping, et al., Deep learning techniques in intelligent fault diagnosis and prognosis for industrial systems: A review, Sensors 23 (3) (2023) 1305. https://doi.org/10.3390/s23031305.',
    '[21] H. Ma, Z. Zhang, J. Zhao, A novel ST-YOLO network for steel-surface-defect detection, Sensors 23 (22) (2023) 9152. https://doi.org/10.3390/s23229152.',
    '[22] K. Ragab, DSTEELNet: A real-time parallel dilated CNN with atrous spatial pyramid pooling for detecting and classifying defects in surface steel strips, Sensors 23 (1) (2023) 544. https://doi.org/10.3390/s23010544.',
    '[23] J. Brodeur, R. Pellerin, I. Deschamps, Operationalization of critical success factors to manage the Industry 4.0 transformation of manufacturing SMEs, Sustainability 14 (14) (2022) 8954. https://doi.org/10.3390/su14148954.',
    '[24] A. Bianchini, I. Savini, A. Andreoni, M. Morolli, V. Solfrini, Manufacturing execution system application within manufacturing small–medium enterprises towards key performance indicators development and their implementation in the production line, Sustainability 16 (7) (2024) 2974. https://doi.org/10.3390/su16072974.',
    '[25] Q. Chen, W. Zhang, N. Jin, X. Wang, P. Dai, Digital transformation evaluation for small- and medium-sized manufacturing enterprises using the fuzzy synthetic method DEMATEL-ANP, Sustainability 14 (20) (2022) 13038. https://doi.org/10.3390/su142013038.',
    '[26] M.-A. Roy, G. Abdul-Nour, S. Gamache, Implementation of an Industry 4.0 strategy adapted to manufacturing SMEs: Simulation and case study, Sustainability 15 (21) (2023) 15423. https://doi.org/10.3390/su152115423.',
    '[27] Y. Zhou, A. Păcală, Empowering manufacturing SMEs: Financial accessibility and sustainable practices in the age of digitalization, Sustainability 17 (8) (2025) 3571. https://doi.org/10.3390/su17083571.',
]
for reference in references:
    p=doc.add_paragraph(reference)
    p.paragraph_format.left_indent=Inches(.23)
    p.paragraph_format.first_line_indent=Inches(-.23)
    p.paragraph_format.line_spacing=2.0
    p.paragraph_format.space_after=Pt(0)
    for run in p.runs: run.font.size=Pt(9)

# KBS uses numbered in-text citations and an appearance-ordered bibliography.
# The source bibliography above is grouped by publisher/topic, so renumber it
# *after* composition. This prevents invalid numeric mappings when sections move.
cite_pattern = re.compile(r'\[([0-9,\s\-–]+)\]')
ref_heading = next(i for i, p in enumerate(doc.paragraphs) if p.text == 'References')
body_paragraphs = doc.paragraphs[:ref_heading]
reference_paragraphs = doc.paragraphs[ref_heading + 1:]

def expand_citation(label):
    numbers = []
    for part in label.split(','):
        part = part.strip()
        if '–' in part or '-' in part:
            bounds = re.split(r'\s*[–-]\s*', part)
            if len(bounds) != 2:
                raise ValueError(f'Unrecognized citation range: {label}')
            numbers.extend(range(int(bounds[0]), int(bounds[1]) + 1))
        elif part:
            numbers.append(int(part))
    return numbers

first_use = []
for paragraph in body_paragraphs:
    for match in cite_pattern.finditer(paragraph.text):
        for source_id in expand_citation(match.group(1)):
            if source_id not in first_use:
                first_use.append(source_id)

assert set(first_use) == set(range(1, len(references) + 1)), 'Uncited or invalid reference number'
number_map = {old: new for new, old in enumerate(first_use, start=1)}

def format_new_citations(match):
    numbers = sorted({number_map[n] for n in expand_citation(match.group(1))})
    components = []
    i = 0
    while i < len(numbers):
        j = i
        while j + 1 < len(numbers) and numbers[j + 1] == numbers[j] + 1:
            j += 1
        if j - i >= 2:
            components.append(f'{numbers[i]}–{numbers[j]}')
        else:
            components.extend(str(n) for n in numbers[i:j + 1])
        i = j + 1
    return '[' + ','.join(components) + ']'

for paragraph in body_paragraphs:
    if cite_pattern.search(paragraph.text):
        # All citation-bearing body paragraphs are generated as one text run.
        assert len(paragraph.runs) == 1, 'Split citation runs need explicit handling'
        paragraph.runs[0].text = cite_pattern.sub(format_new_citations, paragraph.text)

by_original_number = {index: value for index, value in enumerate(references, start=1)}
assert len(reference_paragraphs) == len(references)
for new_index, (old_index, paragraph) in enumerate(zip(first_use, reference_paragraphs), start=1):
    paragraph.text = re.sub(r'^\[\d+\]', f'[{new_index}]', by_original_number[old_index])
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9)

out=BASE/('KBS_Expanded_Research_Manuscript.docx' if EXPANDED else 'manuscript.docx')
doc.save(out)
print(out)
